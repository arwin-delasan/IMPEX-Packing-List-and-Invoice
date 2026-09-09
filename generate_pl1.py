#!/usr/bin/env python3
"""Generate a categorized, SO-style packing list (matching the historical
"PL1" format) from a SAWO packing-list PDF.

Unlike extract_packing_list.py (which fills a fixed-layout template),
this format has a variable number of rows depending on how many
category/subtype groups a shipment has, so there's no blank template to
fill in place. Instead, row styles (fonts, fills, number formats) are
copied from a reference file - the historical "PL1 - Copy.xlsx" example -
onto freshly generated rows.

Categorization comes from product code + name rules (categorize.py), not
Odoo's product category field, which turned out to be inconsistently
maintained for the same product lines. Any product code the rules can't
classify is placed in an "(Uncategorized)" bucket and flagged with a
warning rather than guessed at.

Usage:
    python generate_pl1.py <input.pdf> <output.xlsx>
"""

import copy
import math
import os
import re
import sys
import zipfile
from collections import defaultdict

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.styles import Border

from pdf_parser import ExtractionError
from packing_list_pdf import extract_pdf
from categorize import UNCATEGORIZED, category_sort_key, classify, classify_subgroup, subtype_sort_key
from overrides import OverridesError, get_override, load_overrides
from odoo_client import OdooError, prompt_login
from app_paths import app_path

REFERENCE_PATH = app_path("PL1 - Copy.xlsx")
SHEET_NAME = "summary"

# Column layout (1-indexed), same in the reference and our output.
COL_CODE, COL_NAME, COL_UM, COL_PKG, COL_NET, COL_GROSS, COL_VOL, COL_QTY = range(1, 9)

# Reference rows to copy styles from (found by inspecting PL1 - Copy.xlsx).
REF_ITEM_ROW = 17
REF_CATEGORY_ROW = 16
REF_SUBTYPE_ROW = 26
REF_SUBTOTAL_ROW = 24
REF_GRANDTOTAL_ROW = 97
REF_UNITS_ROW = 98
REF_HEADER_ROW = 14  # "Product Code" / "Product Name" / ... row

# Value cells in the reference file that hold data specific to that one
# historical shipment (Reference Number, Mode of Shipment, Freight Terms,
# Container No., Seal No., Contact Person/No./email) with no PDF or Odoo
# source yet - copy their style but not their stale value.
BLANK_VALUE_CELLS = {(3, 7), (5, 7), (7, 7), (8, 7), (9, 7), (10, 2), (11, 2), (12, 2)}

ROW_HEIGHT_PER_LINE = 15  # points; Excel's default single-line row height


class GenerationError(Exception):
    """Raised when the PDF, Odoo lookup, or reference file doesn't match
    what the generator expects."""


def fail(msg):
    raise GenerationError(msg)


def copy_style(src_cell, dst_cell):
    dst_cell.font = copy.copy(src_cell.font)
    dst_cell.fill = copy.copy(src_cell.fill)
    dst_cell.border = copy.copy(src_cell.border)
    dst_cell.alignment = copy.copy(src_cell.alignment)
    dst_cell.number_format = src_cell.number_format


def derive_so_numbers(header):
    """PDF mislabels its SO numbers as "Invoice No." and "Seal No." -
    combine and clean them into PL1's "SO:" field, e.g. "7119, 7202, 7273"."""
    parts = []
    for raw in (header["seal_no"], header["invoice_no"]):
        for token in raw.split(","):
            token = re.sub(r"(?i)^\s*SO\s*", "", token).strip()
            if token:
                parts.append(token)
    return ", ".join(parts)


def um_warnings(items):
    """Flag rows whose UM code wasn't recognized (pdf_parser.py leaves it
    blank rather than failing the whole extraction) so it doesn't silently
    ship with an empty UM cell."""
    return [
        f"{it['external_code']}: unrecognized UM code on this row - left blank, please fill in manually"
        for it in items if not it["um"]
    ]


def classify_items(items, odoo, overrides=None):
    if overrides is None:
        overrides = load_overrides()

    codes = [it["external_code"] for it in items]
    try:
        products = odoo.lookup_products_by_code(codes)
    except OdooError as e:
        fail(f"Odoo lookup failed: {e}")

    warnings = []
    classified = []
    for it in items:
        code = it["external_code"]
        rec = products.get(code)
        if not rec:
            warnings.append(f"{code}: not found in Odoo (classification may be less reliable)")

        override = get_override(code, overrides)
        if override:
            # A human already corrected this exact code - that always wins
            # over classify()'s guess, no matter what it would otherwise say.
            category, subtype = override["category"], override["subtype"]
        else:
            # The PDF's own description is always used as the displayed name -
            # it's what was actually printed/shipped, and won't drift if Odoo's
            # product name changes later. Odoo's name/category are used only to
            # help classify the item, not to relabel it. Both wordings are fed
            # to classify()'s keyword matching, since Odoo's name and the PDF's
            # description sometimes describe the same product differently
            # (e.g. Odoo's "Wooden Door with Window" vs the PDF's "...with
            # Small Glass Window") - checking only one risks missing a keyword
            # that only appears in the other.
            odoo_name = rec["name"] if rec else it["description"]
            classify_name = odoo_name if odoo_name == it["description"] else f"{odoo_name} {it['description']}"
            categ_path = rec["categ_id"][1] if rec and rec["categ_id"] else None
            category, subtype = classify(code, classify_name, categ_path)
            if category is None:
                category = UNCATEGORIZED
                warnings.append(f"{code} ({it['description']}): no classification rule matched")

        subgroup = classify_subgroup(category, subtype, code, it["description"])
        classified.append({
            **it, "name": it["description"], "category": category, "subtype": subtype,
            "subgroup": subgroup, "overridden": override is not None,
        })
    return classified, warnings


def build_blocks(classified_items):
    """Group items into (category, subtype) blocks for header purposes,
    each holding an ordered list of (subgroup, items) sub-chunks - normally
    just one sub-chunk per block, but classify_subgroup() can split a
    single subtype header into multiple Sub-Totals (see "Wire / Cable")."""
    blocks_map = defaultdict(list)
    block_order = []
    for it in classified_items:
        key = (it["category"], it["subtype"])
        if key not in blocks_map:
            block_order.append(key)
        blocks_map[key].append(it)

    ordered_keys = sorted(
        block_order,
        key=lambda k: (category_sort_key(k[0]), subtype_sort_key(k[0], k[1]), block_order.index(k)),
    )

    blocks = []
    for categ, sub in ordered_keys:
        items = blocks_map[(categ, sub)]
        subchunks_map = defaultdict(list)
        subchunk_order = []
        for it in items:
            if it["subgroup"] not in subchunks_map:
                subchunk_order.append(it["subgroup"])
            subchunks_map[it["subgroup"]].append(it)
        # The "default" sub-chunk (subgroup == the subtype itself, e.g. plain
        # wire under "Wire / Cable") always sorts first; any split-off
        # subgroup (e.g. "Wire / Cable :: Splitter") sorts after it,
        # regardless of which happened to appear first in the PDF.
        first_seen = {sg: i for i, sg in enumerate(subchunk_order)}
        subchunk_order.sort(key=lambda sg: (sg != sub, first_seen[sg]))
        subchunks = [(sg, subchunks_map[sg]) for sg in subchunk_order]
        blocks.append((categ, sub, subchunks))
    return blocks


_SIGNATURE_DRAWING_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<xdr:oneCellAnchor><xdr:from><xdr:col>{from_col}</xdr:col><xdr:colOff>{from_col_off}</xdr:colOff><xdr:row>{from_row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="{width_emu}" cy="{height_emu}"/><xdr:sp><xdr:nvSpPr><xdr:cNvPr id="2" name="Signature Box"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill><a:ln w="6360"><a:solidFill>
<a:srgbClr val="000000"/></a:solidFill><a:miter/></a:ln></xdr:spPr><xdr:style><a:lnRef idx="0"/><a:fillRef idx="0"/><a:effectRef idx="0"/><a:fontRef idx="minor"/></xdr:style><xdr:txBody><a:bodyPr lIns="20160" rIns="20160" tIns="20160" bIns="20160" anchor="t"><a:noAutofit/></a:bodyPr>
<a:p><a:r><a:rPr b="1" lang="en-US" sz="1100"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:latin typeface="Tahoma"/></a:rPr><a:t>S A W O,  I N C.</a:t></a:r></a:p>
<a:p><a:endParaRPr lang="en-US" sz="800"/></a:p><a:p><a:endParaRPr lang="en-US" sz="800"/></a:p><a:p><a:endParaRPr lang="en-US" sz="800"/></a:p><a:p><a:r><a:rPr sz="800"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:latin typeface="Arial Black"/></a:rPr><a:t>{underline}</a:t></a:r></a:p><a:p><a:r><a:rPr sz="1000"><a:solidFill><a:srgbClr val="000000"/></a:solidFill>
<a:latin typeface="Tahoma"/></a:rPr><a:t>  Authorized Signatory</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>"""

_UNDERSCORE_EMU_PER_CHAR = 62388  # calibrated: 36 underscores looked right at a 2286286-EMU-wide box
_TEXT_BODY_MARGIN_EMU = 20160 * 2  # bodyPr's lIns + rIns above


def _col_width_emu(ws, col_letter):
    """A column's width in EMU, using the standard approximation for the
    default font (Calibri 11, Maximum Digit Width 7px):
    pixels = round(width * 7 + 5); 1 px = 9525 EMU at 96 DPI."""
    dim = ws.column_dimensions.get(col_letter)
    width = dim.width if dim and dim.width else 8.43  # Excel's own default column width
    return round(width * 7 + 5) * 9525


def _emu_size_for_cell_span(ws, first_col_letter, last_col_letter, first_row, num_rows):
    """Approximate, in EMU, how big a cell range currently renders as -
    used to give the signature shape an explicit, cell-geometry-independent
    size (see _inject_signature_shape) derived from what its old
    cell-stretched size used to be, rather than a made-up absolute number.
    Row height is already in points; 1 pt = 12700 EMU exactly."""
    total_width_emu = sum(
        _col_width_emu(ws, chr(c))
        for c in range(ord(first_col_letter), ord(last_col_letter) + 1)
    )
    total_height_pt = 0.0
    for r in range(first_row, first_row + num_rows):
        dim = ws.row_dimensions.get(r)
        total_height_pt += dim.height if dim and dim.height else 15.0  # Excel's default row height
    return total_width_emu, round(total_height_pt * 12700)


def _x_emu_to_anchor(ws, target_x_emu, last_col_letter="Z"):
    """Convert an absolute X offset (EMU, measured from column A's left
    edge) into an OOXML drawing anchor (0-indexed column, colOff EMU
    within that column) - walks columns left to right accumulating width
    until target_x_emu lands inside one. Used to horizontally position
    the signature shape by an absolute offset (e.g. centered against the
    table) rather than "start of column B"."""
    col_idx = 0  # 0-indexed; column A
    remaining = max(0, round(target_x_emu))
    while True:
        letter = get_column_letter(col_idx + 1)
        width_emu = _col_width_emu(ws, letter)
        if remaining < width_emu or letter == last_col_letter:
            return col_idx, remaining
        remaining -= width_emu
        col_idx += 1


def add_signature_block(ws, output_path, sig_anchor_row, table_last_col):
    """Size, center, and inject the "S A W O, I N C. / Authorized
    Signatory" floating shape - shared by both generate_pl1.py and
    generate_inv.py so the two documents' signature boxes stay visually
    identical. Must be called after the workbook has already been saved
    to output_path (the shape is added by patching the saved .xlsx's
    OOXML directly - see _inject_signature_shape). sig_anchor_row is the
    0-indexed row the shape's top edge sits on; the caller is expected to
    have reserved 4 blank rows there. table_last_col is the last column
    letter of the document's own item table (e.g. "H" for the packing
    list, "F" for the invoice) - used to center the shape horizontally
    against that table specifically, not the wider print area.
    """
    # Explicit size: 70% of the width / 140% of the height it used to
    # render as when it was just stretched to fit a reserved B:D x 4-row
    # cell span (see _emu_size_for_cell_span), then width narrowed by
    # another 10% on top of that, then height bumped another 10% on top
    # of the 140% (the extra blank paragraph in the drawing XML puts that
    # growth in the gap between the title and the underline, not below
    # "Authorized Signatory", so the underline's own position relative to
    # the signature line beneath it doesn't shift).
    current_width_emu, current_height_emu = _emu_size_for_cell_span(ws, "B", "D", sig_anchor_row, 4)
    sig_width_emu = round(current_width_emu * 0.70 * 0.90)
    sig_height_emu = round(current_height_emu * 1.40 * 1.10)

    # Center the shape horizontally against the document's actual item
    # table, not the wider print area (which can include blank
    # template-leftover columns).
    last_col_idx = ord(table_last_col.upper()) - ord("A") + 1
    table_width_emu = sum(_col_width_emu(ws, get_column_letter(c)) for c in range(1, last_col_idx + 1))
    target_x_emu = max(0, (table_width_emu - sig_width_emu) / 2)
    sig_from_col, sig_from_col_off = _x_emu_to_anchor(ws, target_x_emu)

    # Underline spans the shape's full inner width (shape width minus the
    # text body's left/right insets), rather than a fixed guessed count.
    underline_chars = max(1, round((sig_width_emu - _TEXT_BODY_MARGIN_EMU) / _UNDERSCORE_EMU_PER_CHAR))
    sig_underline = "_" * underline_chars

    _inject_signature_shape(
        output_path, sig_anchor_row, sig_from_col, sig_from_col_off,
        sig_width_emu, sig_height_emu, sig_underline,
    )


def _inject_signature_shape(output_path, anchor_row, from_col, from_col_off, width_emu, height_emu, underline_text):
    """Add the "S A W O, I N C. / Authorized Signatory" box as a real
    floating shape, matching the reference file's actual drawing XML
    (confirmed by inspecting it directly - it's a "Rectangle 2" shape, not
    cell content). openpyxl has no API for writing arbitrary shapes, so
    this patches the .xlsx openpyxl just saved: read it back as a zip, add
    the drawing part plus its relationship/content-type wiring, and
    rewrite the archive. anchor_row is 0-indexed. Anchored via
    oneCellAnchor + an explicit ext (width_emu/height_emu), not stretched
    to fit a cell range, so its size stays exactly what's asked for
    regardless of column widths/row heights.
    """
    drawing_xml = _SIGNATURE_DRAWING_XML.format(
        from_col=from_col, from_col_off=from_col_off, from_row=anchor_row,
        width_emu=width_emu, height_emu=height_emu, underline=underline_text,
    )

    sheet_part = "xl/worksheets/sheet1.xml"  # the only sheet -> openpyxl always names it sheet1.xml
    sheet_rels_part = "xl/worksheets/_rels/sheet1.xml.rels"
    drawing_part = "xl/drawings/drawing1.xml"

    with zipfile.ZipFile(output_path, "r") as zin:
        names = zin.namelist()
        contents = {n: zin.read(n) for n in names}

    rels_xml = contents.get(sheet_rels_part)
    if rels_xml is None:
        rels_xml = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b"</Relationships>"
        )
        names.append(sheet_rels_part)
    rels_text = rels_xml.decode("utf-8")
    existing_ids = [int(i) for i in re.findall(r'Id="rId(\d+)"', rels_text)]
    drawing_rid = f"rId{max(existing_ids, default=0) + 1}"
    rels_text = rels_text.replace(
        "</Relationships>",
        f'<Relationship Id="{drawing_rid}" '
        f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
        f'Target="../drawings/drawing1.xml"/></Relationships>',
    )
    contents[sheet_rels_part] = rels_text.encode("utf-8")

    sheet_text = contents[sheet_part].decode("utf-8")
    if 'xmlns:r=' not in sheet_text.split(">", 1)[0]:
        # openpyxl doesn't declare the relationships namespace prefix on
        # <worksheet> at all (it has no r:-prefixed attributes of its own),
        # so <drawing r:id="..."> below would reference an undeclared
        # namespace - silently ignored by strict consumers - without this.
        sheet_text = sheet_text.replace(
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
            1,
        )
    sheet_text = sheet_text.replace("</worksheet>", f'<drawing r:id="{drawing_rid}"/></worksheet>')
    contents[sheet_part] = sheet_text.encode("utf-8")

    contents[drawing_part] = drawing_xml.encode("utf-8")
    if drawing_part not in names:
        names.append(drawing_part)

    ct_text = contents["[Content_Types].xml"].decode("utf-8")
    ct_text = ct_text.replace(
        "</Types>",
        '<Override PartName="/xl/drawings/drawing1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>',
    )
    contents["[Content_Types].xml"] = ct_text.encode("utf-8")

    tmp_path = output_path + ".tmp"
    with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in names:
            zout.writestr(name, contents[name])
    os.replace(tmp_path, output_path)


def build_workbook(header, blocks, output_path):
    ref_wb = load_workbook(REFERENCE_PATH)
    ref_ws = ref_wb[SHEET_NAME]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    # --- copy the static title + header-field block (rows 1-13) verbatim ---
    for r in range(1, REF_HEADER_ROW):
        for c in range(1, ref_ws.max_column + 1):
            src = ref_ws.cell(row=r, column=c)
            if (r, c) in BLANK_VALUE_CELLS:
                copy_style(src, ws.cell(row=r, column=c))
            elif src.value not in (None, ""):
                dst = ws.cell(row=r, column=c, value=src.value)
                copy_style(src, dst)
    for merge_range in ref_ws.merged_cells.ranges:
        if merge_range.max_row < REF_HEADER_ROW:
            ws.merge_cells(str(merge_range))
    for col_letter, dim in ref_ws.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[col_letter].width = dim.width
    ws.row_dimensions[REF_HEADER_ROW - 1].height = 6  # spacer row directly above the column header

    # --- overwrite the header fields we can actually source ---
    ws.cell(row=3, column=COL_NAME, value=header["sold_to_company"])
    for i, line in enumerate(header["sold_to_address"], start=1):
        ws.cell(row=3 + i, column=COL_NAME, value=line)
    ws.cell(row=4, column=7, value=header["date"])  # already MM/DD/YYYY via extract_pdf
    ws.cell(row=6, column=7, value=header["destination"])
    ws.cell(row=10, column=7, value=derive_so_numbers(header))

    # --- copy the item-table header row (Product Code / Name / UM / ...) ---
    for c in range(1, COL_QTY + 1):
        src = ref_ws.cell(row=REF_HEADER_ROW, column=c)
        dst = ws.cell(row=REF_HEADER_ROW, column=c, value=src.value)
        copy_style(src, dst)

    def style_row(dst_row, ref_row, values=None):
        """Copy a reference row's style across ALL 8 columns (not just the
        ones with visible text) so borders/fills on empty cells carry over
        too, then overlay any given {col: value}."""
        values = values or {}
        for c in range(1, COL_QTY + 1):
            cell = ws.cell(row=dst_row, column=c, value=values.get(c))
            copy_style(ref_ws.cell(row=ref_row, column=c), cell)

    name_col_width = ws.column_dimensions[get_column_letter(COL_NAME)].width or 34

    row = REF_HEADER_ROW + 1
    style_row(row, REF_HEADER_ROW + 1)  # blank spacer row between the column header and the first block
    ws.row_dimensions[row].height = 3
    # No border on this row's left/right outer edges specifically.
    left_cell = ws.cell(row=row, column=COL_CODE)
    left_cell.border = Border(top=left_cell.border.top, bottom=left_cell.border.bottom, right=left_cell.border.right)
    right_cell = ws.cell(row=row, column=COL_QTY)
    right_cell.border = Border(top=right_cell.border.top, bottom=right_cell.border.bottom, left=right_cell.border.left)
    row += 1
    subtotal_rows = []
    prev_category = None

    for category, subtype, subchunks in blocks:
        if category != prev_category:
            style_row(row, REF_CATEGORY_ROW, {COL_NAME: f"***{category}***"})
            row += 1
            prev_category = category

        if subtype:
            style_row(row, REF_SUBTYPE_ROW, {COL_CODE: subtype})
            row += 1

        # Normally one sub-chunk per header; classify_subgroup() can split
        # a single subtype header into several Sub-Totals (see "Wire /
        # Cable" splitting wire from splitters, or "Sauna Heater" splitting
        # off steam generators). A split-off subgroup is keyed as
        # "{subtype} :: {label}" - print that trailing label as its own
        # sub-header (same style as the subtype row) so it reads the same
        # way "Sauna Heater" itself does, rather than silently blending
        # into the parent subtype's rows.
        for _subgroup, chunk_items in subchunks:
            if _subgroup and " :: " in _subgroup:
                sub_label = _subgroup.split(" :: ", 1)[1]
                style_row(row, REF_SUBTYPE_ROW, {COL_CODE: sub_label})
                row += 1
            first_item_row = row
            for it in chunk_items:
                vals = {
                    COL_CODE: it["external_code"], COL_NAME: it["name"], COL_UM: it["um"],
                    COL_PKG: it["packages"], COL_NET: it["net_weight"],
                    COL_GROSS: it["gross_weight"], COL_VOL: it["cbm"], COL_QTY: it["qty"],
                }
                style_row(row, REF_ITEM_ROW, vals)
                line_count = max(1, math.ceil(len(it["name"]) / name_col_width))
                if line_count > 1:
                    ws.row_dimensions[row].height = ROW_HEIGHT_PER_LINE * line_count
                row += 1
            last_item_row = row - 1

            subtotal_vals = {COL_NAME: "Sub-Total:"}
            for c in (COL_PKG, COL_NET, COL_GROSS, COL_VOL, COL_QTY):
                col_letter = get_column_letter(c)
                subtotal_vals[c] = f"=SUM({col_letter}{first_item_row}:{col_letter}{last_item_row})"
            style_row(row, REF_SUBTOTAL_ROW, subtotal_vals)
            subtotal_rows.append(row)
            row += 1

    # --- "Nothing Follows" marker, boxed spacer, then Grand Total + units ---
    style_row(row, 95, {COL_NAME: "*** Nothing Follows ***"})
    row += 1

    style_row(row, 96)  # blank spacer row - carries the box's top border
    row += 1

    gt_row = row
    gt_vals = {COL_NAME: "GRAND TOTAL:"}
    for c in (COL_PKG, COL_NET, COL_GROSS, COL_VOL, COL_QTY):
        col_letter = get_column_letter(c)
        refs = ",".join(f"{col_letter}{r}" for r in subtotal_rows)
        gt_vals[c] = f"=SUM({refs})"
    style_row(row, REF_GRANDTOTAL_ROW, gt_vals)
    row += 1

    units_row = row
    units_vals = dict(zip((COL_PKG, COL_NET, COL_GROSS, COL_VOL, COL_QTY),
                           ("Crate", "kgs.", "kgs.", "cbm.", "qty")))
    style_row(row, REF_UNITS_ROW, units_vals)

    # "GRAND TOTAL:" spans both rows vertically, like the reference's B97:B98
    ws.merge_cells(start_row=gt_row, start_column=COL_NAME, end_row=units_row, end_column=COL_NAME)
    ws.merge_cells(start_row=gt_row, start_column=COL_UM, end_row=units_row, end_column=COL_UM)
    row += 2

    # --- signature block ---
    # The reference's "SAWO, INC. / Authorized Signatory" box is a real
    # floating shape (confirmed via its raw drawing XML), not cell content.
    # openpyxl has no API for writing arbitrary shapes, so it's added by
    # injecting the OOXML drawing part directly after saving (see
    # _inject_signature_shape). Reserve 4 blank rows here for it to sit
    # over, matching its height in the reference.
    sig_anchor_row = row
    row += 4

    # Print area, matching the reference ('A1:L102' there) - this is also
    # what makes Excel draw the blue print-area boundary box in Normal view.
    ws.print_area = f"A1:L{row}"
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=3.6, bottom=1.3, header=1.1, footer=0.6)

    # Print Titles: repeat the title/header block (rows 1-15) on every
    # printed page. Without this, Excel has no reason to keep page breaks
    # aligned with the content's logical sections, and multi-page printouts
    # jump between top/bottom and left/right pages in a confusing order
    # (reported by user as a jumbled Page Break Preview).
    ws.print_title_rows = "1:15"

    # Rest of the page setup, copied verbatim from the reference's own
    # sheet1.xml (paperSize=9 is A4). fitToPage=False means Excel uses
    # `scale`, not fitToWidth/fitToHeight, to shrink the sheet to fit -
    # matches the reference exactly.
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.scale = 95
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "portrait"
    ws.page_setup.pageOrder = "downThenOver"
    ws.page_setup.horizontalDpi = 300
    ws.page_setup.verticalDpi = 300
    ws.print_options.headings = False
    ws.print_options.gridLines = False
    ws.oddHeader.right.text = "Page &P of  &N"

    wb.save(output_path)
    add_signature_block(ws, output_path, sig_anchor_row, get_column_letter(COL_QTY))


def validate(items, grand_total):
    def close(a, b, tol):
        return abs(a - b) <= tol

    sums = {
        "Net Weight": (sum(it["net_weight"] for it in items), grand_total["net_weight"], 0.05),
        "Gross Weight": (sum(it["gross_weight"] for it in items), grand_total["gross_weight"], 0.05),
        "CBM": (sum(it["cbm"] for it in items), grand_total["cbm"], 0.05),
        "Qty": (sum(it["qty"] for it in items), grand_total["qty"], 0),
        "No. of Packages": (
            sum(it["packages"] for it in items if isinstance(it["packages"], int)),
            grand_total["packages"], 0,
        ),
    }
    warnings = []
    for name, (extracted, expected, tol) in sums.items():
        if not close(extracted, expected, tol):
            warnings.append(f"{name} mismatch — extracted sum {extracted} vs PDF Grand Total {expected}")
    return warnings


def main():
    argv = list(sys.argv[1:])
    odoo_pdf = "--odoo" in argv
    if odoo_pdf:
        argv.remove("--odoo")
    if len(argv) != 2:
        print("Usage: python generate_pl1.py [--odoo] <input.pdf> <output.xlsx>", file=sys.stderr)
        sys.exit(1)

    pdf_path, output_path = argv

    try:
        # --odoo still forces that template; without it the file itself
        # decides (see packing_list_pdf).
        header, items, grand_total, template = extract_pdf(
            pdf_path, template="Odoo PDF" if odoo_pdf else None
        )
        print(f"✅ Detected template: {template}.")
        print(f"✅ Extracted header block and {len(items)} line items from PDF.")
        for w in um_warnings(items):
            print(f"WARNING: {w}")

        odoo = prompt_login()
        print("✅ Connected to Odoo.")

        classified, warnings = classify_items(items, odoo)
        for w in warnings:
            print(f"WARNING: {w}")
        print(f"✅ Classified {len(classified)} items ({len(warnings)} warning(s)).")

        blocks = build_blocks(classified)
        print(f"✅ Grouped into {len(blocks)} category/subtype block(s).")

        build_workbook(header, blocks, output_path)
        print(f"✅ Wrote output workbook: {output_path}")
    except (ExtractionError, GenerationError, OdooError, OverridesError) as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)

    totals_warnings = validate(items, grand_total)
    if not totals_warnings:
        print("✅ Validation passed: extracted totals reconcile with PDF Grand Total.")
    else:
        for w in totals_warnings:
            print(f"WARNING: {w}")
        print("Validation FAILED — see WARNING(s) above.", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
